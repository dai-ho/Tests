"""Import du marquage des temps (heures saisies par agent et par élément récepteur).

Fichier attendu (Excel .xlsx ou CSV), 4 colonnes dans cet ordre :
  1. Genre Prénom NOM   (ex. « M. Jean-Pierre DUPONT », « Mme Marie DE LA FONTAINE »)
  2. Objet récepteur    (centre de coût, ex. « C-TOTO_/10 »)
  3. date               (jour, ou mois : 15/03/2024, 03/2024, 03.2024, 2024-03, date Excel…)
  4. Somme de Heures saisies

Les lignes d'en-tête, de total et les lignes vides sont ignorées ; dans un export de tableau
croisé, un nom ou un code laissé vide reprend la valeur de la ligne précédente.
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from models import Agent, Parametre, TempsSaisi, db

HEURES_PAR_ETP_DEFAUT = 1607  # durée annuelle légale du travail
CIVILITES = {"m", "mr", "monsieur", "mme", "madame", "mlle", "mademoiselle", "dr", "pr"}
# Un profil créé par l'import sans temps saisi depuis ce délai (avant la fin du fichier)
# est considéré comme parti à la date de sa dernière saisie
DELAI_DEPART = timedelta(days=183)


def heures_par_etp():
    p = Parametre.query.get("heures_par_etp")
    try:
        return float(p.valeur) if p and p.valeur else HEURES_PAR_ETP_DEFAUT
    except ValueError:
        return HEURES_PAR_ETP_DEFAUT


def definir_heures_par_etp(valeur):
    db.session.merge(Parametre(cle="heures_par_etp", valeur=str(valeur)))
    db.session.commit()


# --- Lecture du fichier ----------------------------------------------------------

def normaliser(texte):
    texte = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texte.lower()).strip()


def decouper_nom(brut):
    """« M. Jean-Pierre DUPONT » -> ("Jean-Pierre", "DUPONT") : le prénom précède le
    premier mot entièrement en majuscules, le nom de famille est la suite."""
    mots = str(brut).split()
    while mots and normaliser(mots[0]).replace(" ", "") in CIVILITES:
        mots = mots[1:]
    if not mots:
        return None, None
    for i, mot in enumerate(mots):
        lettres = [c for c in mot if c.isalpha()]
        if i > 0 and lettres and all(c.isupper() for c in lettres):
            return " ".join(mots[:i]), " ".join(mots[i:])
    if len(mots) == 1:
        return None, mots[0]
    return " ".join(mots[:-1]), mots[-1]


def lire_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # numéro de série Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    m = re.fullmatch(r"(\d{1,2})[/.\-](\d{4})", s) or re.fullmatch(r"(\d{4})[/.\-](\d{1,2})", s)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        annee, mois = (b, a) if a <= 12 else (a, b)
        if 1 <= mois <= 12:
            return date(annee, mois, 1)
    return None


def lire_heures(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(s) if s else None
    except ValueError:
        return None


def lignes_brutes(fichier, nom_fichier):
    """Les 4 premières colonnes de chaque ligne du fichier (Excel ou CSV)."""
    if nom_fichier.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(fichier, data_only=True, read_only=True)
        for ligne in wb.worksheets[0].iter_rows(values_only=True):
            yield list(ligne[:4]) + [None] * (4 - len(ligne[:4]))
        return
    contenu = fichier.read()
    texte = None
    for codage in ("utf-8-sig", "cp1252"):
        try:
            texte = contenu.decode(codage)
            break
        except UnicodeDecodeError:
            pass
    try:
        separateur = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        separateur = ";"  # cas usuel des exports français ; un fichier illisible sera signalé
    for ligne in csv.reader(io.StringIO(texte), delimiter=separateur):
        yield (ligne[:4] + [None] * 4)[:4]


def lire_fichier(fichier, nom_fichier):
    """Renvoie (saisies, ignorees) : saisies = [(nom brut, code, date, heures)]."""
    saisies, ignorees = [], []
    nom_courant = code_courant = None
    for num, (nom, code, jour, heures) in enumerate(lignes_brutes(fichier, nom_fichier), start=1):
        textes = [str(x).strip() for x in (nom, code, jour) if x not in (None, "")]
        if not textes and heures in (None, ""):
            continue
        if any(t.lower().startswith(("total", "somme", "étiquette", "etiquette")) for t in textes):
            continue
        if nom not in (None, ""):
            nom_courant = str(nom).strip()
        if code not in (None, ""):
            code_courant = str(code).strip()
        d, h = lire_date(jour), lire_heures(heures)
        if d is None or h is None or not nom_courant or not code_courant:
            if num > 1 or saisies:  # la 1re ligne illisible est l'en-tête
                ignorees.append(num)
            continue
        saisies.append((nom_courant, code_courant, d, h))
    return saisies, ignorees


# --- Import ------------------------------------------------------------------------

def importer(fichier, nom_fichier):
    """Met à jour les temps saisis de la période couverte par le fichier et crée les
    profils manquants. Renvoie un compte rendu (dict)."""
    saisies, ignorees = lire_fichier(fichier, nom_fichier)
    if not saisies:
        raise ValueError("aucune ligne exploitable (4 colonnes attendues : Genre Prénom NOM, "
                         "objet récepteur, date, heures)")
    debut, fin = min(s[2] for s in saisies), max(s[2] for s in saisies)

    agents = {}
    for a in Agent.query:
        agents[(normaliser(a.prenom), normaliser(a.nom))] = a
        agents.setdefault(("", normaliser(a.nom)), a)
    crees = []
    # Agrégation : plusieurs lignes identiques (même agent, code, date) sont additionnées
    cumul = {}
    for nom_brut, code, d, h in saisies:
        prenom, nom = decouper_nom(nom_brut)
        agent = agents.get((normaliser(prenom), normaliser(nom)))
        if agent is None:
            agent = Agent(nom=nom, prenom=prenom, etp=1.0,
                          commentaire=f"Fiche créée par l'import des temps ({nom_brut}).")
            db.session.add(agent)
            agents[(normaliser(prenom), normaliser(nom))] = agent
            crees.append(agent)
        cumul[(agent, code, d)] = cumul.get((agent, code, d), 0) + h
    db.session.flush()

    # Remplacement des temps de la période couverte par le fichier
    supprimes = TempsSaisi.query.filter(TempsSaisi.date >= debut, TempsSaisi.date <= fin).delete()
    for (agent, code, d), h in cumul.items():
        db.session.add(TempsSaisi(agent_id=agent.id, code=code, date=d, heures=h))

    # Profils créés : arrivée à la 1re saisie ; partis si plus de saisie depuis longtemps
    partis = []
    for agent in crees:
        dates = [d for (a, _, d) in cumul if a is agent]
        agent.date_arrivee = min(dates)
        if max(dates) < fin - DELAI_DEPART:
            agent.date_depart = max(dates)
            partis.append(agent)
    db.session.commit()
    return {"lignes": len(saisies), "enregistrements": len(cumul), "remplaces": supprimes,
            "debut": debut, "fin": fin, "ignorees": ignorees,
            "heures": sum(cumul.values()), "crees": crees, "partis": partis,
            "agents": len({a for (a, _, _) in cumul})}


# --- Synthèses -----------------------------------------------------------------------

def heures_par(annee, cle):
    """Heures de l'année regroupées selon `cle` : "agent", "code" ou "agent_code"."""
    debut, fin = date(annee, 1, 1), date(annee, 12, 31)
    resultat = {}
    for t in TempsSaisi.query.filter(TempsSaisi.date >= debut, TempsSaisi.date <= fin):
        k = {"agent": t.agent_id, "code": t.code, "agent_code": (t.agent_id, t.code)}[cle]
        resultat[k] = resultat.get(k, 0) + t.heures
    return resultat


def codes_connus():
    """Centres de coût / éléments récepteurs connus : temps importés, DA, budgets, affectations."""
    from formulaires import valeurs_distinctes
    from models import Affectation, LigneBudgetaire, PosteDA
    return sorted(set(valeurs_distinctes(TempsSaisi.code))
                  | set(valeurs_distinctes(PosteDA.code_imputation))
                  | set(valeurs_distinctes(LigneBudgetaire.code))
                  | set(valeurs_distinctes(Affectation.projet)))
