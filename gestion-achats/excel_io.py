"""Import / export Excel (.xlsx) au format du fichier « Budgets GT »
(onglets « Suivi DA » et « Suivi BdC », plus « Marchés » ajouté par l'application)."""
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models import (TYPES_BUDGET, BudgetAnnuel, Commande, DemandeAchat, LigneBudgetaire, Marche,
                    PosteCommande, PosteDA, db)

# Colonnes reconnues : clé interne -> en-têtes acceptés (le premier sert à l'export).
# Les en-têtes sont comparés sans accents, sans casse ni ponctuation.
# Les colonnes calculées d'Excel (DA@poste, Imputé, Disponible) sont recalculées par l'application.
COLONNES_DA = {
    "annee": ["Année", "Annee budgétaire"],
    "projet": ["Projet"],
    "reference": ["DA", "N° DA", "Num DA", "Numéro DA", "Référence DA"],
    "numero": ["poste", "N° poste", "Poste DA"],
    "nom": ["Désignation", "Nom", "Libellé", "Intitulé"],
    "montant": ["Valeur", "Montant", "Montant HT"],
    "code_imputation": ["Centre de coût", "Code d'imputation", "Imputation"],
    "budget": ["Budget"],
    "statut": ["Statut"],
    "remarque": ["Remarque", "Remarques", "Commentaire"],
    "date_livraison": ["Date de livraison", "Date livraison"],
    "demandeur": ["Demandeur"],
    "certificateur": ["Certificateur"],
}

COLONNES_CMD = {
    "societe": ["Société", "Fournisseur"],
    "activite": ["Activité"],
    "da_designation": ["Imputation", "Désignation DA"],
    "reference": ["BDC/Marché", "BDC", "N° BDC", "N° Commande", "Commande"],
    "marche": ["N° marché", "Marché", "Numéro de marché"],
    "lot": ["Lot"],
    "numero": ["poste du BDC", "Poste", "Poste commande"],
    "nom": ["Désignation", "Nom", "Libellé"],
    "montant": ["Montant"],
    "date_livraison": ["Echéance", "Date de livraison", "Date livraison"],
    "da_cle": ["DA@poste"],
    "ptf": ["PAF", "PTF"],
    "csf": ["CSF"],
    "ref_livrable": ["CR pour CSF", "Référence livrable", "Livrable"],
}

COLONNES_MARCHE = {
    "reference": ["Marché", "N° marché", "Référence"],
    "titulaire": ["Titulaire", "Fournisseur", "Société"],
    "lots": ["Lots", "Lot"],
    "simplifie": ["Simplifié", "Marché simplifié"],
    "date_echeance": ["Échéance", "Date d'échéance", "Fin"],
    "montant_max": ["Montant maximal", "Montant max", "Maximum", "Plafond"],
    "commentaire": ["Commentaire", "Remarque"],
}

COLONNES_BUDGET = {
    "code": ["Code", "Ligne budgétaire", "Centre de coût"],
    "libelle": ["Libellé", "Intitulé"],
    "type": ["Type"],
    "projet": ["Projet"],
    "annee": ["Année"],
    "montant_demande": ["Demandé", "Montant demandé"],
    "montant_alloue": ["Alloué", "Montant alloué"],
    "report": ["Report"],
    "commentaire_ligne": ["Commentaire ligne"],
    "commentaire": ["Commentaire année", "Commentaire"],
}

VRAI = {"oui", "o", "x", "1", "fait", "ok", "yes", "y", "vrai", "true", "certifie"}


def _normaliser(texte):
    texte = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texte.lower()).strip()


def _texte(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _entier(v):
    s = _texte(v)
    return int(float(s.replace(",", "."))) if s else None


def _montant(v):
    if v is None or isinstance(v, (int, float)):
        return v
    s = re.sub(r"[^\d,.\-]", "", str(v)).replace(",", ".")
    return float(s) if s else None


def _date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"date illisible « {s} »")


def _booleen(v):
    if isinstance(v, bool):
        return v
    return _normaliser(v) in VRAI


def _statut(v):
    s = _normaliser(v)
    if s.startswith("supprim"):
        return "supprimé"
    if s.startswith("clos"):
        return "clos"
    return _texte(v) or ""


def _lire_feuille(ws, colonnes, obligatoires):
    """Renvoie la liste des lignes (dict clé interne -> valeur) d'une feuille."""
    alias = {}
    for cle, noms in colonnes.items():
        for a in noms:
            alias.setdefault(_normaliser(a), cle)
    lignes = ws.iter_rows(values_only=True)
    for num_entete, ligne in enumerate(lignes, start=1):
        mapping = {}
        for j, cellule in enumerate(ligne):
            cle = alias.get(_normaliser(cellule))
            if cle and cle not in mapping.values():
                mapping[j] = cle
        if obligatoires <= set(mapping.values()):
            break
        if num_entete >= 15:
            raise ValueError(f"onglet « {ws.title} » : ligne d'en-tête introuvable "
                             f"(colonnes attendues : {', '.join(colonnes[c][0] for c in obligatoires)})")
    else:
        raise ValueError(f"onglet « {ws.title} » vide")

    resultat = []
    for num_ligne, ligne in enumerate(lignes, start=num_entete + 1):
        valeurs = {cle: ligne[j] if j < len(ligne) else None for j, cle in mapping.items()}
        if all(_texte(v) is None for v in valeurs.values()):
            continue
        valeurs["_ligne"] = num_ligne
        resultat.append(valeurs)
    return resultat


def _trouver_feuilles(wb):
    feuille_da = feuille_cmd = feuille_marches = feuille_budgets = None
    for ws in wb.worksheets:
        mots = set(_normaliser(ws.title).split())
        if feuille_cmd is None and ({"bdc", "bc", "commandes", "commande"} & mots):
            feuille_cmd = ws
        elif feuille_marches is None and ({"marche", "marches"} & mots):
            feuille_marches = ws
        elif feuille_budgets is None and ({"budget", "budgets"} & mots):
            feuille_budgets = ws
        elif feuille_da is None and ({"da", "demande", "demandes"} & mots):
            feuille_da = ws
    return feuille_da, feuille_cmd, feuille_marches, feuille_budgets


def _vider(avec_referentiels):
    PosteCommande.query.delete()
    Commande.query.delete()
    if avec_referentiels:
        Marche.query.delete()
        BudgetAnnuel.query.delete()
        LigneBudgetaire.query.delete()
    PosteDA.query.delete()
    DemandeAchat.query.delete()
    db.session.flush()


def _importer_da(feuille, erreurs):
    das = {d.reference: d for d in DemandeAchat.query.filter(DemandeAchat.reference.isnot(None))}
    # Postes sans n° de DA : retrouvés par leur désignation lors d'une mise à jour
    sans_numero = {p.nom: p for p in PosteDA.query.join(DemandeAchat)
                   .filter(DemandeAchat.reference.is_(None))}
    nb = 0
    for v in _lire_feuille(feuille, COLONNES_DA, {"reference", "numero", "nom"}):
        ou = f"{feuille.title} ligne {v['_ligne']}"
        try:
            ref, nom = _texte(v.get("reference")), _texte(v.get("nom"))
            numero = _entier(v.get("numero"))
            if not nom:
                if ref or numero:
                    raise ValueError("désignation manquante")
                continue  # ligne ne contenant que des valeurs annexes
            # Conversion complète avant toute écriture : une ligne invalide est ignorée
            champs = dict(
                nom=nom,
                montant=_montant(v.get("montant")),
                annee=_entier(v.get("annee")),
                projet=_texte(v.get("projet")),
                code_imputation=_texte(v.get("code_imputation")),
                budget=_montant(v.get("budget")),
                statut=_statut(v.get("statut")),
                remarque=_texte(v.get("remarque")),
                date_livraison=_date(v.get("date_livraison")),
                demandeur=_texte(v.get("demandeur")),
                certificateur=_texte(v.get("certificateur")),
            )
            if ref:
                if numero is None:
                    raise ValueError("numéro de poste manquant")
                da = das.get(ref)
                if da is None:
                    da = das[ref] = DemandeAchat(reference=ref)
                    db.session.add(da)
                poste = next((p for p in da.postes if p.numero == numero), None)
                if poste is None:
                    poste = PosteDA(numero=numero)
                    da.postes.append(poste)
            else:
                poste = sans_numero.get(nom)
                if poste is None:
                    da = DemandeAchat()
                    db.session.add(da)
                    poste = sans_numero[nom] = PosteDA(numero=numero or 10)
                    da.postes.append(poste)
            for cle, valeur in champs.items():
                setattr(poste, cle, valeur)
            nb += 1
        except (ValueError, TypeError) as e:
            erreurs.append(f"{ou} : {e}")
    db.session.flush()
    return nb


def _importer_marches(feuille, erreurs):
    marches = {m.reference: m for m in Marche.query}
    nb = 0
    for v in _lire_feuille(feuille, COLONNES_MARCHE, {"reference"}):
        try:
            ref = _texte(v.get("reference"))
            if not ref:
                continue
            champs = dict(titulaire=_texte(v.get("titulaire")), lots=_texte(v.get("lots")),
                          simplifie=_booleen(v.get("simplifie")),
                          date_echeance=_date(v.get("date_echeance")),
                          montant_max=_montant(v.get("montant_max")),
                          commentaire=_texte(v.get("commentaire")))
            marche = marches.get(ref)
            if marche is None:
                marche = marches[ref] = Marche(reference=ref)
                db.session.add(marche)
            for cle, valeur in champs.items():
                setattr(marche, cle, valeur)
            nb += 1
        except (ValueError, TypeError) as e:
            erreurs.append(f"{feuille.title} ligne {v['_ligne']} : {e}")
    db.session.flush()
    return nb


def _importer_budgets(feuille, erreurs):
    lignes = {l.code: l for l in LigneBudgetaire.query}
    for v in _lire_feuille(feuille, COLONNES_BUDGET, {"code"}):
        try:
            code = _texte(v.get("code"))
            if not code:
                continue
            annee = _entier(v.get("annee"))
            montants = {c: _montant(v.get(c)) for c in ("montant_demande", "montant_alloue", "report")}
            ligne = lignes.get(code)
            if ligne is None:
                ligne = lignes[code] = LigneBudgetaire(code=code)
                db.session.add(ligne)
            type_ = _normaliser(v.get("type"))
            ligne.type = "projet" if type_.startswith("projet") else "structure"
            ligne.libelle = _texte(v.get("libelle")) or ligne.libelle
            ligne.projet = _texte(v.get("projet")) or ligne.projet
            ligne.commentaire = _texte(v.get("commentaire_ligne")) or ligne.commentaire
            if annee:
                b = ligne.budget(annee)
                if b is None:
                    b = BudgetAnnuel(annee=annee)
                    ligne.annees.append(b)
                for c, m in montants.items():
                    setattr(b, c, m)
                b.commentaire = _texte(v.get("commentaire"))
        except (ValueError, TypeError) as e:
            erreurs.append(f"{feuille.title} ligne {v['_ligne']} : {e}")
    db.session.flush()


def _trouver_poste_da(v, par_designation, par_cle):
    """Retrouve le poste de DA d'une ligne de BdC : par la désignation (colonne
    « Imputation », comme la RECHERCHEV d'Excel), sinon par « DA@poste »."""
    designation = _texte(v.get("da_designation"))
    candidats = par_designation.get(_normaliser(designation), []) if designation else []
    cle = _texte(v.get("da_cle"))
    m = re.match(r"\s*([^@\s]*)@(\d+)", cle or "")
    if m:
        poste = par_cle.get((m.group(1), int(m.group(2))))
        if poste and (not candidats or poste in candidats):
            return poste, None
    if len(candidats) == 1:
        return candidats[0], None
    if len(candidats) > 1:
        return None, f"désignation de DA « {designation} » en double, poste importé sans lien"
    if designation or cle:
        return None, f"poste de DA « {designation or cle} » introuvable, poste importé sans lien"
    return None, None


def _importer_cmd(feuille, erreurs):
    par_designation, par_cle = {}, {}
    for p in PosteDA.query.all():
        par_designation.setdefault(_normaliser(p.nom), []).append(p)
        if p.da.reference:
            par_cle[(p.da.reference, p.numero)] = p
    cmds = {c.reference: c for c in Commande.query.filter(Commande.reference.isnot(None))}
    # Postes de BdC sans n° : retrouvés par leur désignation lors d'une mise à jour
    sans_numero = {p.nom: p for p in PosteCommande.query.join(Commande)
                   .filter(Commande.reference.is_(None))}
    marches = {m.reference: m for m in Marche.query}
    nb = 0
    for v in _lire_feuille(feuille, COLONNES_CMD, {"reference", "numero"}):
        ou = f"{feuille.title} ligne {v['_ligne']}"
        try:
            ref, numero = _texte(v.get("reference")), _entier(v.get("numero"))
            if not ref and not (_texte(v.get("nom")) or _texte(v.get("da_designation"))):
                continue  # ligne vide ou ne contenant que des formules
            if ref and numero is None:
                raise ValueError("n° de poste manquant")
            csf = _booleen(v.get("csf"))
            poste_da, alerte = _trouver_poste_da(v, par_designation, par_cle)
            champs = dict(
                nom=_texte(v.get("nom")) or (poste_da.nom if poste_da else f"Poste {numero}"),
                societe=_texte(v.get("societe")),
                activite=_texte(v.get("activite")),
                montant=_montant(v.get("montant")),
                date_livraison=_date(v.get("date_livraison")),
                ptf=_texte(v.get("ptf")),
                csf=csf,
                # La référence du livrable justifie la CSF : ignorée si la CSF n'est pas faite
                ref_livrable=_texte(v.get("ref_livrable")) if csf else None,
                poste_da=poste_da,
            )
            ref_marche = _texte(v.get("marche"))
            marche = marches.get(ref_marche) if ref_marche else None
            if ref_marche and marche is None:
                # Marché inconnu : créé avec sa seule référence, à compléter dans l'onglet Marchés
                marche = marches[ref_marche] = Marche(reference=ref_marche,
                                                      titulaire=champs["societe"])
                db.session.add(marche)
                erreurs.append(f"{ou} : marché {ref_marche} créé, à compléter "
                               "(montant maximal, échéance)")
            if ref:
                cmd = cmds.get(ref)
                if cmd is None:
                    cmd = cmds[ref] = Commande(reference=ref)
                    db.session.add(cmd)
                poste = next((p for p in cmd.postes if p.numero == numero), None)
                if poste is None:
                    poste = PosteCommande(numero=numero)
                    cmd.postes.append(poste)
            else:
                # BdC dont le n° n'est pas encore connu : un BdC par ligne
                poste = sans_numero.get(champs["nom"])
                if poste is None:
                    cmd = Commande()
                    db.session.add(cmd)
                    poste = sans_numero[champs["nom"]] = PosteCommande(numero=numero or 10)
                    cmd.postes.append(poste)
                cmd = poste.commande
            if ref_marche:
                cmd.marche = marche
            if _texte(v.get("lot")):
                cmd.lot = _texte(v.get("lot"))
            for cle, valeur in champs.items():
                setattr(poste, cle, valeur)
            if alerte:
                erreurs.append(f"{ou} : {alerte}")
            nb += 1
        except (ValueError, TypeError) as e:
            erreurs.append(f"{ou} : {e}")
    return nb


def importer(fichier, vider=False, avec_referentiels=False):
    """Importe un classeur : onglets DA et BdC, plus les onglets Marchés et Budgets si
    `avec_referentiels` (fichier exporté par l'application). Les autres onglets sont ignorés.
    Renvoie (nb_postes_da, nb_postes_cmd, erreurs, onglets importés)."""
    wb = load_workbook(fichier, data_only=True, read_only=True)
    feuille_da, feuille_cmd, feuille_marches, feuille_budgets = _trouver_feuilles(wb)
    if feuille_da is None and feuille_cmd is None:
        raise ValueError("aucun onglet « Suivi DA » ou « Suivi BdC » trouvé")
    if not avec_referentiels:
        feuille_marches = feuille_budgets = None
    if vider:
        # Marchés et budgets saisis dans l'application sont conservés, sauf s'ils sont réimportés
        _vider(avec_referentiels=avec_referentiels)
    erreurs = []
    if feuille_marches is not None:
        _importer_marches(feuille_marches, erreurs)
    if feuille_budgets is not None:
        _importer_budgets(feuille_budgets, erreurs)
    nb_da = _importer_da(feuille_da, erreurs) if feuille_da is not None else 0
    nb_cmd = _importer_cmd(feuille_cmd, erreurs) if feuille_cmd is not None else 0
    db.session.commit()
    onglets = [f.title for f in (feuille_da, feuille_cmd, feuille_marches, feuille_budgets)
               if f is not None]
    return nb_da, nb_cmd, erreurs, onglets


# --- Export ----------------------------------------------------------------------

ENTETES_DA = ["Année", "Projet", "DA", "poste", "Désignation", "DA@poste", "Valeur",
              "Centre de coût", "Budget", "Imputé", "Disponible", "Statut", "Remarque",
              "Date de livraison", "Demandeur", "Certificateur"]
ENTETES_CMD = ["Société", "Activité", "Imputation", "BDC/Marché", "N° marché", "Lot", "poste du BDC",
               "Désignation",
               "Montant", "Echéance", "DA@poste", "PAF", "CSF", "CR pour CSF"]
ENTETES_MARCHE = ["Marché", "Titulaire", "Lots", "Simplifié", "Échéance", "Montant maximal", "Commandé", "Reste",
                  "Commentaire"]
ENTETES_BUDGET = ["Code", "Libellé", "Type", "Projet", "Année", "Demandé", "Alloué", "Report",
                  "Commentaire année", "Commentaire ligne"]
COLONNES_EUROS = {"Valeur", "Budget", "Imputé", "Disponible", "Montant", "Montant maximal",
                  "Commandé", "Reste", "Demandé", "Alloué", "Report"}


def _entier_si_possible(s):
    return int(s) if s and s.isdigit() else s


def _ecrire_feuille(ws, entetes, lignes):
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="305496")
    for ligne in lignes:
        ws.append(ligne)
    for i, entete in enumerate(entetes, start=1):
        lettre = get_column_letter(i)
        ws.column_dimensions[lettre].width = 40 if entete in ("Désignation", "Remarque", "Imputation") \
            else max(12, len(entete) + 4)
        for cellule in ws[lettre][1:]:
            if isinstance(cellule.value, date):
                cellule.number_format = "DD/MM/YYYY"
            elif entete in COLONNES_EUROS:
                cellule.number_format = '#,##0.00\\ "€"'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def exporter():
    """Exporte toute la base au format du fichier Excel (réimportable)."""
    postes_da = (PosteDA.query.join(DemandeAchat)
                 .order_by(DemandeAchat.reference.is_(None), DemandeAchat.reference,
                           PosteDA.numero).all())
    postes_cmd = (PosteCommande.query.join(Commande)
                  .order_by(Commande.reference.isnot(None), Commande.reference,
                            PosteCommande.numero).all())
    wb = Workbook()
    ws_da = wb.active
    ws_da.title = "Suivi DA"
    _ecrire_feuille(ws_da, ENTETES_DA, [
        [p.annee, p.projet, _entier_si_possible(p.da.reference), p.numero, p.nom, p.cle,
         p.montant, p.code_imputation, p.budget, p.impute, p.disponible, p.statut or None,
         p.remarque, p.date_livraison, p.demandeur, p.certificateur]
        for p in postes_da])
    _ecrire_feuille(wb.create_sheet("Suivi BdC"), ENTETES_CMD, [
        [p.societe, p.activite, p.poste_da.nom if p.poste_da else None,
         _entier_si_possible(p.commande.reference),
         p.commande.marche.reference if p.commande.marche else None, p.commande.lot, p.numero, p.nom, p.montant, p.date_livraison,
         p.poste_da.cle if p.poste_da else None, p.ptf, "OK" if p.csf else None, p.ref_livrable]
        for p in postes_cmd])
    _ecrire_feuille(wb.create_sheet("Marchés"), ENTETES_MARCHE, [
        [m.reference, m.titulaire, m.lots, "Oui" if m.simplifie else None, m.date_echeance, m.montant_max, m.engage(), m.restant,
         m.commentaire]
        for m in Marche.query.order_by(Marche.reference)])
    lignes_budget = []
    for l in LigneBudgetaire.query.order_by(LigneBudgetaire.code):
        debut = [l.code, l.libelle, l.type, l.projet]
        for b in l.annees or [None]:
            lignes_budget.append(debut + (
                [b.annee, b.montant_demande, b.montant_alloue, b.report, b.commentaire] if b
                else [None] * 5) + [l.commentaire])
    _ecrire_feuille(wb.create_sheet("Budgets"), ENTETES_BUDGET, lignes_budget)
    flux = BytesIO()
    wb.save(flux)
    flux.seek(0)
    return flux
